"""Official-file-only historical IPO import.

This deliberately does not call KRX, KIND, KIS, or DART. Historical data is
accepted only from an operator-uploaded official KRX/KIND export and is first
held in an opaque, user-bound preview before it can be committed.
"""
from __future__ import annotations

import csv
import hashlib
from html.parser import HTMLParser
import io
import json
import math
import re
import secrets
import threading
from copy import deepcopy
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook

from app.services.ipo.identity import generate_ipo_id, is_spac_ipo, normalize_company_name
from app.services.ipo.store import _STORE_LOCK, _read_market_store_unlocked, _write_market_store_unlocked

MAX_UPLOAD_BYTES = 16 * 1024 * 1024
PREVIEW_TTL_SECONDS = 15 * 60
MAX_SERVER_PREVIEWS = 128
_HEADER_SCAN_ROWS = 20

_PREVIEWS: dict[str, dict[str, Any]] = {}
_PREVIEW_LOCK = threading.RLock()

SOURCE_KRX = "KRX_NEW_LISTINGS_EXPORT"
SOURCE_KIND = "KIND_NEW_LISTING_COMPANIES_EXPORT"

ALIASES = {
    "company_name": {"종목명", "회사명", "법인명", "기업명"},
    "stock_code": {"종목코드", "단축코드", "주식종목코드"},
    "actual_listing_date": {"상장일", "신규상장일", "상장일자"},
    "final_offer_price": {"공모가", "공모가(원)", "공모가격", "확정공모가", "공모가격(원)"},
    "offering_amount": {
        "공모금액",
        "공모총액",
        "공모금액(원)",
        "공모금액(천원)",
        "공모금액(백만원)",
        "공모금액(억원)",
    },
    "market": {"시장구분", "시장", "소속시장", "시장명"},
    "shares": {"공모주식수", "공모주수", "공모수량"},
    "listing_type": {"상장유형"},
    "lead_manager": {
        "상장주선인",
        "상장주선인(지정자문인)",
        "상장주선인/지정자문인",
        "대표주관회사",
        "주관사",
        "지정자문인",
    },
}

# These signatures are intentionally stricter than the field aliases. They are
# used only to identify an official export family, not to infer arbitrary files.
_KRX_SOURCE_MARKERS = {
    "증권구분",
    "주식종류",
    "상장유형",
    "상장주선인(지정자문인)",
    "공모주식수",
    "최초상장주식수",
    "현재상장주식수",
}
_KIND_SOURCE_MARKERS = {
    "신규상장일",
    "대표주관회사",
    "주요제품",
    "공모금액",
    "공모금액(원)",
    "공모금액(천원)",
    "공모금액(백만원)",
    "공모금액(억원)",
    "상장주선인/지정자문인",
    "지정자문인",
}


class HistoricalImportError(ValueError):
    def __init__(self, code: str, message: str = "공식 과거 공모주 파일을 확인할 수 없습니다."):
        super().__init__(message)
        self.code = code


def _header(value: object) -> str:
    return str(value or "").replace(" ", "").replace("\n", "").replace("\r", "").strip().lower()


def _digest(value: object) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def _market_digest(market: dict[str, Any]) -> str:
    """Digest canonical market content without timestamp-only false staleness."""
    return _digest(
        {
            "schema_version": market.get("schema_version"),
            "ipos": market.get("ipos") if isinstance(market.get("ipos"), list) else [],
        }
    )


def _safe_filename(filename: str | None) -> str:
    return Path(str(filename or "official-export")).name[:180]


def _current_kst_date() -> date:
    try:
        tz = ZoneInfo("Asia/Seoul")
    except ZoneInfoNotFoundError:
        tz = timezone(timedelta(hours=9), name="Asia/Seoul")
    return datetime.now(tz).date()


def _mapping_for_headers(headers: list[str]) -> dict[str, int]:
    normalized = [_header(x) for x in headers]
    mapping: dict[str, int] = {}
    for field, aliases in ALIASES.items():
        aliases_n = {_header(alias) for alias in aliases}
        matches = [index for index, value in enumerate(normalized) if value in aliases_n]
        if len(matches) > 1:
            raise HistoricalImportError("SOURCE_UNRECOGNIZED", f"'{field}' 열을 하나로 식별할 수 없습니다.")
        if matches:
            mapping[field] = matches[0]
    return mapping


def _find_header_row(rows: list[list[Any]]) -> tuple[list[str], list[list[Any]], int]:
    required = {"company_name", "stock_code", "actual_listing_date", "final_offer_price"}
    for index, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        headers = [str(value or "") for value in row]
        try:
            mapping = _mapping_for_headers(headers)
        except HistoricalImportError:
            continue
        if required.issubset(mapping):
            # Excel/CSV row numbers are 1-based, so the first data row follows
            # the detected header row.
            return headers, rows[index + 1 :], index + 2
    raise HistoricalImportError(
        "SOURCE_UNRECOGNIZED",
        "공식 KRX 또는 KIND 신규상장 파일의 필수 열을 찾을 수 없습니다.",
    )


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise HistoricalImportError("MALFORMED_WORKBOOK", "CSV 문자 인코딩을 확인할 수 없습니다.")


class _HtmlTableParser(HTMLParser):
    """Minimal parser for KIND's official HTML-disguised-as-XLS export."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self._table_depth = 0
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        tag = tag.lower()
        if tag == "table":
            self._table_depth += 1
        elif self._table_depth and tag == "tr":
            self._row = []
        elif self._row is not None and tag in {"th", "td"}:
            self._cell = []
        elif self._cell is not None and tag == "br":
            self._cell.append(" ")

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"th", "td"} and self._cell is not None:
            value = re.sub(r"\s+", " ", "".join(self._cell)).strip()
            if self._row is not None:
                self._row.append(value)
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if any(str(value).strip() for value in self._row):
                self.rows.append(self._row)
            self._row = None
            self._cell = None
        elif tag == "table" and self._table_depth:
            self._table_depth -= 1


def _decode_kind_html(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise HistoricalImportError(
        "MALFORMED_WORKBOOK",
        "KIND XLS 문자 인코딩을 확인할 수 없습니다.",
    )


def _read_kind_html_rows(content: bytes) -> list[list[str]]:
    text = _decode_kind_html(content)
    head = text[:65536].lower()
    if (
        "<html" not in head
        or "<table" not in head
        or "신규상장기업현황" not in text[:65536]
    ):
        raise HistoricalImportError(
            "MALFORMED_WORKBOOK",
            "KIND 신규상장기업현황 XLS 형식을 확인할 수 없습니다.",
        )

    parser = _HtmlTableParser()
    parser.feed(text)
    parser.close()
    if not parser.rows:
        raise HistoricalImportError("MALFORMED_WORKBOOK")
    return parser.rows


def _read_rows(filename: str, content: bytes) -> tuple[list[str], list[list[Any]], int]:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xlsm", ".xls", ".csv"}:
        raise HistoricalImportError(
            "UNSUPPORTED_FILE_TYPE",
            "xlsx, xlsm, xls 또는 CSV 파일만 가져올 수 있습니다.",
        )
    if not content:
        raise HistoricalImportError("EMPTY_FILE", "비어 있는 파일은 가져올 수 없습니다.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HistoricalImportError("FILE_TOO_LARGE", "파일 크기가 16MB 제한을 초과했습니다.")

    try:
        if suffix == ".csv":
            rows = list(csv.reader(io.StringIO(_decode_csv(content))))
        elif suffix == ".xls":
            rows = _read_kind_html_rows(content)
        else:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            try:
                sheet = workbook.active
                # Some official KRX exports contain an incorrect worksheet
                # dimension ("A1") even though the XML contains the full table.
                # openpyxl read-only mode trusts that metadata unless reset.
                if sheet.calculate_dimension() == "A1:A1":
                    sheet.reset_dimensions()
                rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            finally:
                workbook.close()
    except HistoricalImportError:
        raise
    except Exception as exc:
        raise HistoricalImportError("MALFORMED_WORKBOOK") from exc

    if not rows:
        raise HistoricalImportError("MALFORMED_WORKBOOK")
    return _find_header_row(rows)


def detect_source(headers: list[str]) -> tuple[str, dict[str, int]]:
    normalized = [_header(x) for x in headers]
    mapping = _mapping_for_headers(headers)
    required = {"company_name", "stock_code", "actual_listing_date", "final_offer_price"}
    if not required.issubset(mapping):
        raise HistoricalImportError(
            "SOURCE_UNRECOGNIZED",
            "공식 KRX 또는 KIND 신규상장 파일의 필수 열을 찾을 수 없습니다.",
        )

    header_set = set(normalized)
    # KRX Data Marketplace [20001] uses "종목명"; KIND 신규상장기업현황
    # uses "회사명". Secondary markers prevent a generic user-created sheet
    # from being accepted only because four business columns happen to match.
    krx = (
        "종목명" in header_set
        and "회사명" not in header_set
        and bool(header_set & {_header(x) for x in _KRX_SOURCE_MARKERS})
    )
    kind = (
        "회사명" in header_set
        and "종목명" not in header_set
        and bool(header_set & {_header(x) for x in _KIND_SOURCE_MARKERS})
    )
    if krx == kind:
        raise HistoricalImportError(
            "SOURCE_UNRECOGNIZED",
            "공식 KRX 또는 KIND 파일 출처를 열 제목만으로 안전하게 식별할 수 없습니다.",
        )
    return (SOURCE_KRX if krx else SOURCE_KIND), mapping


def _stock_code(value: object) -> str | None:
    if isinstance(value, bool):
        return None
    text = str(value or "").strip()
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    if text.isdigit() and 1 <= len(text) <= 6:
        return text.zfill(6)
    return None


def _date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip().replace(".", "-").replace("/", "-")
    if text.isdigit() and len(text) == 8:
        text = f"{text[:4]}-{text[4:6]}-{text[6:]}"
    try:
        parsed = date.fromisoformat(text)
    except ValueError:
        return None
    return parsed.isoformat()


def _decimal(value: object) -> Decimal | None:
    if isinstance(value, bool):
        return None
    text = str(value or "").strip().replace(",", "")
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    return number if number.is_finite() else None


def _positive_number(value: object) -> int | float | None:
    number = _decimal(value)
    if number is None or number <= 0:
        return None
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _same_number(left: object, right: object) -> bool:
    a, b = _decimal(left), _decimal(right)
    return a is not None and b is not None and a == b


def _same_date(left: object, right: object) -> bool:
    a, b = _date(left), _date(right)
    return a is not None and b is not None and a == b


def _market(value: object) -> str | None:
    text = str(value or "").strip().upper()
    if text in {"코스닥", "KSQ", "KOSDAQ"}:
        return "KOSDAQ"
    if text in {"유가증권", "코스피", "STK", "KOSPI"}:
        return "KOSPI"
    if text in {"코넥스", "KNX", "KONEX"}:
        return "KONEX"
    return text or None


def _lead_managers(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    names = [item.strip() for item in re.split(r"[,;/\n\r]+", text) if item.strip()]
    return list(dict.fromkeys(names))


def _offering_amount(value: object, header: str) -> int | float | None:
    number = _positive_number(value)
    if number is None:
        return None
    h = _header(header)
    if "천원" in h:
        return number * 1_000
    if "백만원" in h:
        return number * 1_000_000
    if "억원" in h:
        return number * 100_000_000
    if "(원)" in h:
        return number
    return None


def _candidate(
    row: list[Any],
    headers: list[str],
    mapping: dict[str, int],
    source: str,
    filename: str,
) -> tuple[dict[str, Any] | None, str | None]:
    def value(field: str) -> object:
        return row[mapping[field]] if field in mapping and mapping[field] < len(row) else None

    company = str(value("company_name") or "").strip()
    code = _stock_code(value("stock_code"))
    listed = _date(value("actual_listing_date"))
    price = _positive_number(value("final_offer_price"))

    if not company:
        return None, "COMPANY_NAME_MISSING"
    if not code:
        return None, "STOCK_CODE_INVALID"
    if not listed:
        return None, "ACTUAL_LISTING_DATE_INVALID"
    if date.fromisoformat(listed) > _current_kst_date():
        return None, "ACTUAL_LISTING_DATE_IN_FUTURE"
    if price is None:
        return None, "FINAL_OFFER_PRICE_INVALID"

    listing_type = str(value("listing_type") or "").replace(" ", "").strip()
    if listing_type and listing_type != "신규상장":
        return None, "NOT_NEW_LISTING"

    provenance: dict[str, Any] = {
        "source": source,
        "imported_at": datetime.now().astimezone().isoformat(),
        "source_filename": filename,
        "actual_listing_date": listed,
        "final_offer_price": price,
        "stock_code": code,
    }

    incoming: dict[str, Any] = {
        "ipo_id": generate_ipo_id(company, stock_code=code),
        "company_name": company,
        "stock_code": code,
        "market": _market(value("market")),
        "listing_track": "general",
        "subscription_start": None,
        "subscription_end": None,
        "payment_date": None,
        "refund_date": None,
        "expected_listing_date": None,
        "actual_listing_date": listed,
        "final_offer_price": price,
        "lead_managers": _lead_managers(value("lead_manager")),
        "features": {},
        "score": {},
        "sources": {"official_historical_import": provenance},
    }

    if "offering_amount" in mapping:
        raw_amount = value("offering_amount")
        amount = _offering_amount(raw_amount, headers[mapping["offering_amount"]])
        if amount is not None:
            provenance["offering_amount"] = amount
            incoming["offering_amount"] = amount
        elif str(raw_amount or "").strip():
            # Preserve only a short scalar for audit; never keep the whole raw row.
            provenance["offering_amount_raw"] = str(raw_amount)[:80]

    if "shares" in mapping:
        shares = _positive_number(value("shares"))
        if shares is not None:
            incoming["offer_shares"] = shares

    incoming["sources"]["official_historical_imports"] = {
        source: deepcopy(provenance),
    }
    incoming["listing_track"] = "spac" if is_spac_ipo(incoming) else "general"
    return incoming, None


def _merge_official_historical_provenance(
    existing: dict[str, Any],
    incoming: dict[str, Any],
) -> None:
    # Preserve the first official import and retain later official enrichments by source.
    existing_sources = existing.setdefault("sources", {})
    if not isinstance(existing_sources, dict):
        existing_sources = {}
        existing["sources"] = existing_sources

    incoming_sources = incoming.get("sources")
    if not isinstance(incoming_sources, dict):
        return

    existing_primary = existing_sources.get("official_historical_import")
    incoming_primary = incoming_sources.get("official_historical_import")

    by_source = existing_sources.get("official_historical_imports")
    if not isinstance(by_source, dict):
        by_source = {}

    def remember(provenance: object) -> None:
        if not isinstance(provenance, dict):
            return
        source_name = str(provenance.get("source") or "").strip()
        if not source_name:
            return
        by_source.setdefault(source_name, deepcopy(provenance))

    # Backfill records created before the multi-source provenance structure existed.
    remember(existing_primary)

    incoming_by_source = incoming_sources.get("official_historical_imports")
    if isinstance(incoming_by_source, dict):
        for source_name, provenance in incoming_by_source.items():
            if not isinstance(provenance, dict):
                continue
            key = str(source_name or provenance.get("source") or "").strip()
            if key:
                by_source.setdefault(key, deepcopy(provenance))

    remember(incoming_primary)

    # Keep the original/primary official source stable. Only populate it when absent.
    if not isinstance(existing_primary, dict) and isinstance(incoming_primary, dict):
        existing_sources["official_historical_import"] = deepcopy(incoming_primary)

    if by_source:
        existing_sources["official_historical_imports"] = by_source


def _classify(
    incoming: dict[str, Any],
    existing: list[dict[str, Any]],
) -> tuple[str, str | None, dict[str, Any] | None]:
    matches = [
        item
        for item in existing
        if str(item.get("stock_code") or "").strip() == incoming["stock_code"]
    ]
    if len(matches) > 1:
        return "REVIEW_REQUIRED", "DUPLICATE_STOCK_CODE", None
    if not matches:
        return "NEW", None, None

    current = matches[0]
    if normalize_company_name(str(current.get("company_name") or "")) != normalize_company_name(
        incoming["company_name"]
    ):
        return "CONFLICT", "COMPANY_NAME_CONFLICT", current

    old_listing, new_listing = current.get("actual_listing_date"), incoming.get("actual_listing_date")
    if old_listing not in (None, "") and new_listing not in (None, ""):
        if not _same_date(old_listing, new_listing):
            return "CONFLICT", "ACTUAL_LISTING_DATE_CONFLICT", current

    old_price, new_price = current.get("final_offer_price"), incoming.get("final_offer_price")
    if old_price not in (None, "") and new_price not in (None, ""):
        if not _same_number(old_price, new_price):
            return "CONFLICT", "FINAL_OFFER_PRICE_CONFLICT", current

    enrichment_fields = (
        "actual_listing_date",
        "final_offer_price",
        "market",
        "offering_amount",
        "offer_shares",
    )
    has_blank_enrichment = any(
        current.get(field) in (None, "") and incoming.get(field) not in (None, "")
        for field in enrichment_fields
    )
    current_sources = current.get("sources") if isinstance(current.get("sources"), dict) else {}
    provenance_missing = not isinstance(current_sources.get("official_historical_import"), dict)
    return ("ENRICHABLE" if has_blank_enrichment or provenance_missing else "ALREADY_PRESENT"), None, current


def _merge_compatible_duplicate(base: dict[str, Any], other: dict[str, Any]) -> tuple[bool, str | None]:
    if normalize_company_name(base["company_name"]) != normalize_company_name(other["company_name"]):
        return False, "DUPLICATE_COMPANY_NAME_CONFLICT"
    if not _same_date(base["actual_listing_date"], other["actual_listing_date"]):
        return False, "DUPLICATE_LISTING_DATE_CONFLICT"
    if not _same_number(base["final_offer_price"], other["final_offer_price"]):
        return False, "DUPLICATE_OFFER_PRICE_CONFLICT"
    if base.get("listing_track") != other.get("listing_track"):
        return False, "DUPLICATE_LISTING_TRACK_CONFLICT"

    left_market, right_market = base.get("market"), other.get("market")
    if left_market and right_market and left_market != right_market:
        return False, "DUPLICATE_MARKET_CONFLICT"
    if not left_market and right_market:
        base["market"] = right_market

    for field in ("offering_amount", "offer_shares"):
        left, right = base.get(field), other.get(field)
        if left not in (None, "") and right not in (None, "") and not _same_number(left, right):
            return False, f"DUPLICATE_{field.upper()}_CONFLICT"
        if left in (None, "") and right not in (None, ""):
            base[field] = right

    managers = list(base.get("lead_managers") or [])
    for manager in other.get("lead_managers") or []:
        if manager not in managers:
            managers.append(manager)
    base["lead_managers"] = managers
    return True, None


def _purge_expired_previews_unlocked(now: float) -> None:
    expired = [ticket for ticket, state in _PREVIEWS.items() if float(state.get("expires") or 0) < now]
    for ticket in expired:
        _PREVIEWS.pop(ticket, None)

    if len(_PREVIEWS) >= MAX_SERVER_PREVIEWS:
        oldest = sorted(
            _PREVIEWS,
            key=lambda ticket: float(_PREVIEWS[ticket].get("expires") or 0),
        )
        for ticket in oldest[: max(1, len(_PREVIEWS) - MAX_SERVER_PREVIEWS + 1)]:
            _PREVIEWS.pop(ticket, None)


def create_preview(filename: str, content: bytes, username: str, market: dict[str, Any]) -> dict[str, Any]:
    headers, rows, first_data_row = _read_rows(filename, content)
    source, mapping = detect_source(headers)
    existing = market.get("ipos") if isinstance(market.get("ipos"), list) else []

    issues: list[dict[str, Any]] = []
    parsed_by_code: dict[str, list[tuple[int, dict[str, Any]]]] = {}
    summary = {
        "total_rows": 0,
        "valid_rows": 0,
        "new": 0,
        "enrichable": 0,
        "already_present": 0,
        "conflict": 0,
        "review_required": 0,
        "invalid": 0,
        "duplicate_rows": 0,
    }
    safe_name = _safe_filename(filename)

    for row_number, row in enumerate(rows, first_data_row):
        if not any(str(value or "").strip() for value in row):
            continue
        summary["total_rows"] += 1
        candidate, error = _candidate(row, headers, mapping, source, safe_name)
        if error:
            summary["invalid"] += 1
            issues.append({"row_number": row_number, "reason": error})
            continue
        assert candidate is not None
        summary["valid_rows"] += 1
        parsed_by_code.setdefault(candidate["stock_code"], []).append((row_number, candidate))

    candidates: list[dict[str, Any]] = []
    for stock_code, grouped in parsed_by_code.items():
        first_row, merged = grouped[0][0], deepcopy(grouped[0][1])
        duplicate_conflict: str | None = None

        for _, duplicate in grouped[1:]:
            compatible, reason = _merge_compatible_duplicate(merged, duplicate)
            if not compatible:
                duplicate_conflict = reason or "DUPLICATE_STOCK_CODE_IN_FILE"
                break

        if duplicate_conflict:
            summary["review_required"] += len(grouped)
            for row_number, candidate in grouped:
                issues.append(
                    {
                        "row_number": row_number,
                        "company_name": candidate["company_name"],
                        "stock_code": stock_code,
                        "reason": duplicate_conflict,
                    }
                )
            continue

        if len(grouped) > 1:
            summary["duplicate_rows"] += len(grouped) - 1

        classification, reason, _ = _classify(merged, existing)
        summary[classification.lower()] += 1
        candidates.append(
            {
                "row_number": first_row,
                "classification": classification,
                "reason": reason,
                "record": merged,
            }
        )
        if classification in {"CONFLICT", "REVIEW_REQUIRED"}:
            issues.append(
                {
                    "row_number": first_row,
                    "company_name": merged["company_name"],
                    "stock_code": stock_code,
                    "reason": reason,
                }
            )

    ticket = secrets.token_urlsafe(32)
    now = datetime.now().timestamp()
    with _PREVIEW_LOCK:
        _purge_expired_previews_unlocked(now)
        _PREVIEWS[ticket] = {
            "username": username,
            "expires": now + PREVIEW_TTL_SECONDS,
            "market_digest": _market_digest(market),
            "candidate_digest": _digest(candidates),
            "source": source,
            "candidates": candidates,
            "in_use": False,
        }

    return {
        "source": source,
        "summary": summary,
        "issues": issues[:20],
        "issues_truncated": max(0, len(issues) - 20),
        "preview_ticket": ticket,
    }


def commit_preview(ticket: str, username: str) -> dict[str, Any]:
    token = str(ticket or "")
    now = datetime.now().timestamp()

    with _PREVIEW_LOCK:
        state = _PREVIEWS.get(token)
        if not state or state.get("username") != username:
            _purge_expired_previews_unlocked(now)
            raise HistoricalImportError("PREVIEW_TICKET_INVALID")
        if float(state.get("expires") or 0) < now:
            _PREVIEWS.pop(token, None)
            _purge_expired_previews_unlocked(now)
            raise HistoricalImportError("PREVIEW_TICKET_EXPIRED")
        _purge_expired_previews_unlocked(now)
        if state.get("in_use"):
            raise HistoricalImportError("PREVIEW_TICKET_IN_USE", "이미 반영 중인 미리보기입니다.")
        if _digest(state.get("candidates")) != state.get("candidate_digest"):
            _PREVIEWS.pop(token, None)
            raise HistoricalImportError("PREVIEW_TICKET_INVALID")
        state["in_use"] = True

    try:
        with _STORE_LOCK:
            market = _read_market_store_unlocked()
            if _market_digest(market) != state["market_digest"]:
                with _PREVIEW_LOCK:
                    _PREVIEWS.pop(token, None)
                raise HistoricalImportError(
                    "PREVIEW_STALE",
                    "공모주 데이터가 변경되어 미리보기를 다시 실행해야 합니다.",
                )

            working = deepcopy(market)
            working.setdefault("ipos", [])
            committed_new = enriched = skipped = 0

            for item in state["candidates"]:
                incoming = item["record"]
                classification, _, existing = _classify(incoming, working["ipos"])

                if classification == "NEW":
                    working["ipos"].append(deepcopy(incoming))
                    committed_new += 1
                elif classification == "ENRICHABLE" and existing is not None:
                    for field in (
                        "actual_listing_date",
                        "final_offer_price",
                        "market",
                        "offering_amount",
                        "offer_shares",
                    ):
                        if existing.get(field) in (None, "") and incoming.get(field) not in (None, ""):
                            existing[field] = deepcopy(incoming[field])

                    if not existing.get("lead_managers") and incoming.get("lead_managers"):
                        existing["lead_managers"] = deepcopy(incoming["lead_managers"])

                    _merge_official_historical_provenance(existing, incoming)
                    enriched += 1
                else:
                    skipped += 1

            _write_market_store_unlocked(working)
    except Exception:
        with _PREVIEW_LOCK:
            if token in _PREVIEWS:
                _PREVIEWS[token]["in_use"] = False
        raise

    with _PREVIEW_LOCK:
        _PREVIEWS.pop(token, None)

    return {
        "status": "ok",
        "committed_new": committed_new,
        "enriched": enriched,
        "skipped": skipped,
        "total_after": len(working["ipos"]),
        "market": working,
    }
