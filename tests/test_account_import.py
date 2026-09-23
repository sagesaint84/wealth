from __future__ import annotations

import io
import tempfile
import unittest
from copy import deepcopy
from urllib.parse import unquote
from pathlib import Path
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import app.main as main
from app.services import portfolio
from app.services.broker_holdings_sync import ALL_MARKETS, BrokerHoldingsResult, ProviderHoldingScope


def xlsx(rows, headers=None):
    book = Workbook(); sheet = book.active
    sheet.append(headers or ["소유자", "증권사", "계좌명", "계좌번호", "계좌유형", "원화예수금", "달러예수금", "세액공제적용", "소득구간", "올해연금납입액", "ISA전환입금액", "ISA전환연도"])
    for row in rows: sheet.append(row)
    output = io.BytesIO(); book.save(output); return output.getvalue()


class AccountImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.patch = patch("app.services.portfolio._get_user_dir", return_value=Path(self.temp.name))
        self.patch.start()

    def tearDown(self):
        self.patch.stop(); self.temp.cleanup()

    def test_import_general_account(self):
        result = portfolio.import_account_rows("accounts.xlsx", xlsx([
            ["아빠", "NH투자증권(나무)", "나무", "1234-5678-01", "일반", 1500000, 10, "", "", 0, 0, 2026],
        ]), username="u", allowed_owners={"모두", "아빠", "엄마"})
        self.assertEqual(result["created"], 1)
        general = portfolio.read_portfolio("u")["accounts"][0]
        self.assertEqual(general["account_type"], "general")
        self.assertEqual(general["account_no"], "1234-5678-01")

    def test_import_pension_account(self):
        result = portfolio.import_account_rows("accounts.xlsx", xlsx([
            ["엄마", "삼성증권", "연금", "2234-5678-01", "연금저축", 0, 0, "예", "5500만원이하", 6000000, 0, 2026],
        ]), username="u", allowed_owners={"모두", "아빠", "엄마"})
        self.assertEqual(result["created"], 1)
        pension = portfolio.read_portfolio("u")["accounts"][0]
        self.assertEqual(pension["account_type"], "pension_savings")
        self.assertTrue(pension["tax_deductible"])
        self.assertEqual(pension["yearly_contributions"][0]["deposit"], 6000000)

    def test_import_irp_account(self):
        result = portfolio.import_account_rows("accounts.xlsx", xlsx([
            ["엄마", "미래에셋증권", "IRP", "", "IRP", 0, 0, "아니오", "5500만원초과", 3000000, 0, 2026],
        ]), username="u", allowed_owners={"모두", "아빠", "엄마"})
        account = portfolio.read_portfolio("u")["accounts"][0]
        self.assertEqual(result["created"], 1)
        self.assertEqual(account["account_type"], "irp")
        self.assertFalse(account["tax_deductible"])

    def test_import_isa_account(self):
        result = portfolio.import_account_rows("accounts.xlsx", xlsx([
            ["아빠", "한국투자증권", "ISA", "", "ISA", 0, 0, "", "", 0, 5000000, 2026],
        ]), username="u", allowed_owners={"모두", "아빠", "엄마"})
        account = portfolio.read_portfolio("u")["accounts"][0]
        self.assertEqual(result["created"], 1)
        self.assertEqual(account["account_type"], "isa")
        self.assertEqual(account["isa_transfer_amount"], 5000000)

    def test_import_cash_balances_and_yearly_contribution(self):
        result = portfolio.import_account_rows("accounts.xlsx", xlsx([
            ["아빠", "NH투자증권(나무)", "나무", "1234-5678-01", "일반", 1500000, 10, "", "", 0, 0, 2026],
            ["엄마", "삼성증권", "연금", "2234-5678-01", "연금저축", 0, 0, "예", "5500만원이하", 6000000, 0, 2026],
        ]), username="u", allowed_owners={"모두", "아빠", "엄마"})
        self.assertEqual(result["created"], 2)
        data = portfolio.read_portfolio("u"); accounts = data["accounts"]
        pension = next(a for a in accounts if a["name"] == "연금")
        self.assertEqual(pension["yearly_contributions"][0]["deposit"], 6000000)
        self.assertTrue(pension["tax_deductible"])
        general = next(a for a in accounts if a["name"] == "나무")
        self.assertEqual(data["settings"]["cash_balances"][general["id"]], {"KRW": 1500000.0, "USD": 10.0})
        self.assertEqual(data["holdings"], [])

    def test_import_duplicate_by_account_number_and_format(self):
        headers = ["owner", "broker_name", "account_name", "account_number", "account_type", "cash_krw"]
        content = xlsx([["아빠", "NH", "첫계좌", "1234-5678", "general", 1]], headers)
        first = portfolio.import_account_rows("accounts.xlsx", content, username="u", allowed_owners={"모두", "아빠"})
        second = portfolio.import_account_rows("accounts.xlsx", xlsx([["아빠", "NH", "새 별칭", "12345678", "general", 99]], headers), username="u", allowed_owners={"모두", "아빠"})
        self.assertEqual(first["created"], 1); self.assertEqual(second["duplicates"], 1)
        self.assertEqual(portfolio.normalize_broker_account_no(" 1234-5678 "), "12345678")
        self.assertEqual(portfolio.read_portfolio("u")["accounts"][0]["cash_krw"], 1.0)

    def test_import_same_name_different_account_number(self):
        rows = [["아빠", "삼성증권", "종합계좌", "11111111", "일반", 0], ["아빠", "삼성증권", "종합계좌", "22222222", "일반", 0]]
        result = portfolio.import_account_rows("accounts.xlsx", xlsx(rows, ["소유자", "증권사", "계좌명", "계좌번호", "계좌유형", "원화예수금"]), username="u", allowed_owners={"아빠", "모두"})
        self.assertEqual(result["created"], 2)

    def test_import_same_number_different_owner_is_conflict(self):
        headers = ["소유자", "증권사", "계좌명", "계좌번호", "계좌유형"]
        portfolio.import_account_rows("accounts.xlsx", xlsx([["아빠", "키움", "첫", "12345678", "일반"]], headers), username="u", allowed_owners={"아빠", "자녀", "모두"})
        result = portfolio.import_account_rows("accounts.xlsx", xlsx([["자녀", "키움증권", "둘", "12345678", "일반"]], headers), username="u", allowed_owners={"아빠", "자녀", "모두"})
        self.assertEqual((result["created"], result["invalid"]), (0, 1))

    def test_import_without_account_number_uses_fallback_identity(self):
        headers = ["소유자", "증권사", "계좌명", "계좌유형"]
        portfolio.import_account_rows("accounts.xlsx", xlsx([["아빠", "삼성증권", "종합", "일반"]], headers), username="u", allowed_owners={"아빠", "모두"})
        result = portfolio.import_account_rows("accounts.xlsx", xlsx([["아빠", "삼성증권", "종합", "일반"]], headers), username="u", allowed_owners={"아빠", "모두"})
        self.assertEqual(result["duplicates"], 1)

    def test_import_duplicate_does_not_overwrite_or_modify_holdings(self):
        headers = ["소유자", "증권사", "계좌명", "계좌번호", "계좌유형", "원화예수금"]
        portfolio.import_account_rows("accounts.xlsx", xlsx([["아빠", "삼성증권", "종합", "12345678", "일반", 1]], headers), username="u", allowed_owners={"아빠", "모두"})
        data = portfolio.read_portfolio("u"); data["holdings"] = [{"id": "holding", "account_id": data["accounts"][0]["id"]}]; portfolio.write_portfolio(data, "u")
        result = portfolio.import_account_rows("accounts.xlsx", xlsx([["아빠", "삼성증권", "다른 이름", "1234-5678", "일반", 999]], headers), username="u", allowed_owners={"아빠", "모두"})
        saved = portfolio.read_portfolio("u")
        self.assertEqual(result["duplicates"], 1)
        self.assertEqual(saved["accounts"][0]["cash_krw"], 1.0)
        self.assertEqual([item["id"] for item in saved["holdings"]], ["holding"])

    def test_import_rejects_non_finite_numeric_value(self):
        headers = ["소유자", "증권사", "계좌명", "계좌유형", "원화예수금"]
        result = portfolio.import_account_rows("accounts.xlsx", xlsx([["아빠", "삼성", "종합", "일반", "Infinity"]], headers), username="u", allowed_owners={"아빠", "모두"})
        self.assertEqual((result["created"], result["invalid"]), (0, 1))

    def test_import_invalid_owner_type_and_required_fields(self):
        headers = ["owner", "broker_name", "account_name", "account_type", "cash_krw"]
        for row in (["자녀", "삼성", "x", "general", 0], ["아빠", "삼성", "x", "invalid", 0], ["아빠", "", "x", "general", 0], ["아빠", "삼성", "", "general", 0]):
            with self.subTest(row=row):
                result = portfolio.import_account_rows("accounts.xlsx", xlsx([row], headers), username="u", allowed_owners={"아빠", "모두"})
                self.assertEqual(result["invalid"], 1)

    def test_sample_workbook_has_expected_headers_and_numbers(self):
        sample = Path(__file__).resolve().parents[1] / "data" / "샘플_증권계좌.xlsx"
        sheet = load_workbook(sample, data_only=True).active
        self.assertEqual(sheet.title, "증권계좌")
        self.assertEqual([cell.value for cell in sheet[2]][:5], ["소유자", "증권사", "계좌명", "계좌번호", "계좌유형"])
        self.assertIsInstance(sheet.cell(3, 6).value, (int, float))
        self.assertEqual(sheet.cell(3, 4).value, "01234567-01")
        self.assertEqual(sheet.cell(3, 4).number_format, "@")
        self.assertGreaterEqual(sheet.max_row, 6)

    def test_account_card_uses_mask_helper_without_full_number_interpolation(self):
        source = (Path(__file__).resolve().parents[1] / "app" / "static" / "wealth.js").read_text(encoding="utf-8")
        card_region = source[source.index("function renderAccounts"):source.index("function renderHoldings")]
        self.assertIn("function maskBrokerAccountNo", source)
        self.assertIn("계좌 ****${normalized.slice(-4)}", source)
        self.assertIn("maskBrokerAccountNo(account.account_no)", card_region)
        self.assertNotIn("${account.account_no}", card_region)


class AccountImportApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.username = "alice"
        self.paths = patch("app.services.portfolio._get_user_dir", side_effect=lambda username=None: Path(self.temp.name) / str(username or "default"))
        self.paths.start(); self.addCleanup(self.paths.stop)
        self.users = patch("app.services.user_manager.get_user_by_name", side_effect=lambda username: {"username": username, "id": username, "role": "user"})
        self.users.start(); self.addCleanup(self.users.stop)
        self.client = TestClient(main.app)

    def _authenticate(self):
        self.client.cookies.set(main.COOKIE_NAME, main._serializer.dumps({"user": self.username}))

    def _create(self, **overrides):
        self._authenticate()
        payload = {"broker": "NH투자증권(나무)", "account_name": "내 계좌", "owner": "아빠", **overrides}
        return self.client.post("/api/accounts", json=payload)

    def test_post_duplicate_and_same_name_different_number(self):
        first = self._create(account_no="1234-5678-01"); self.assertEqual(first.status_code, 200)
        duplicate = self._create(account_no="1234567801", account_name="다른 이름")
        self.assertEqual(duplicate.status_code, 409); self.assertEqual(duplicate.json()["detail"]["code"], "ACCOUNT_ALREADY_EXISTS")
        second = self._create(account_no="99999999")
        self.assertEqual(second.status_code, 200)
        self.assertEqual(len(portfolio.read_portfolio("alice")["accounts"]), 2)

    def test_post_owner_conflict_and_user_isolation(self):
        self.assertEqual(self._create(account_no="12345678").status_code, 200)
        conflict = self._create(owner="자녀", account_no="12345678")
        self.assertEqual(conflict.status_code, 409); self.assertEqual(conflict.json()["detail"]["code"], "ACCOUNT_NUMBER_OWNER_CONFLICT")
        self.username = "bob"
        self.assertEqual(self._create(account_no="12345678").status_code, 200)
        self.assertEqual(len(portfolio.read_portfolio("alice")["accounts"]), 1)
        self.assertEqual(len(portfolio.read_portfolio("bob")["accounts"]), 1)

    def test_put_account_number_add_change_remove_preserve_and_conflict(self):
        first = self._create(); account_id = first.json()["id"]
        self._authenticate(); self.assertEqual(self.client.put(f"/api/accounts/{account_id}", json={"name": "내 계좌", "account_no": "1234-5678"}).status_code, 200)
        self.assertEqual(portfolio.read_portfolio("alice")["accounts"][0]["account_no"], "1234-5678")
        self.assertEqual(self.client.put(f"/api/accounts/{account_id}", json={"name": "이름만 변경"}).status_code, 200)
        self.assertEqual(portfolio.read_portfolio("alice")["accounts"][0]["account_no"], "1234-5678")
        self.assertEqual(self.client.put(f"/api/accounts/{account_id}", json={"name": "이름만 변경", "account_no": "8765-4321"}).status_code, 200)
        self._create(account_name="다른", account_no="1111-1111")
        conflict = self.client.put(f"/api/accounts/{account_id}", json={"name": "깨지면 안 됨", "account_no": "11111111"})
        self.assertEqual(conflict.status_code, 409)
        self.assertEqual(portfolio.read_portfolio("alice")["accounts"][0]["name"], "이름만 변경")
        self.assertEqual(self.client.put(f"/api/accounts/{account_id}", json={"name": "이름만 변경", "account_no": ""}).status_code, 200)
        self.assertEqual(portfolio.read_portfolio("alice")["accounts"][0]["account_no"], "")

    def test_put_account_number_owner_conflict_does_not_mutate_target(self):
        first = self._create(account_no="12345678"); first_id = first.json()["id"]
        second = self._create(account_name="자녀 계좌", owner="자녀", account_no="87654321"); second_id = second.json()["id"]
        response = self.client.put(f"/api/accounts/{second_id}", json={"name": "바뀌면 안 됨", "account_no": "12345678"})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["detail"]["code"], "ACCOUNT_NUMBER_OWNER_CONFLICT")
        target = next(account for account in portfolio.read_portfolio("alice")["accounts"] if account["id"] == second_id)
        self.assertEqual((target["name"], target["account_no"]), ("자녀 계좌", "87654321"))
        self.assertNotEqual(first_id, second_id)

    def test_import_endpoint_and_sample_download(self):
        body = xlsx([["아빠", "삼성증권", "파일 계좌", "00001234-01", "일반", 100]], ["소유자", "증권사", "계좌명", "계좌번호", "계좌유형", "원화예수금"])
        self._authenticate()
        response = self.client.post("/api/import-accounts", files={"file": ("accounts.xlsx", body, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        self.assertEqual((response.status_code, response.json()["created"]), (200, 1))
        self.username = "bob"; self._authenticate()
        isolated = self.client.post("/api/import-accounts", files={"file": ("accounts.xlsx", body, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        self.assertEqual((isolated.status_code, isolated.json()["created"]), (200, 1))
        sample = self.client.get("/api/sample/accounts")
        self.assertEqual(sample.status_code, 200)
        self.assertIn("샘플_증권계좌.xlsx", unquote(sample.headers["content-disposition"]))


class AccountNumberReconciliationTests(unittest.IsolatedAsyncioTestCase):
    @staticmethod
    def _empty_portfolio(**overrides):
        data = deepcopy(portfolio.EMPTY_PORTFOLIO)
        data.update(overrides)
        data["settings"] = deepcopy(overrides.get("settings", data["settings"]))
        data["settings"].setdefault("cash_balances", {})
        return data

    async def _assert_provider_reuses_manual_account(self, *, provider_attr, provider, runner, account_no, broker, manual_broker=None):
        data = self._empty_portfolio(accounts=[{
            "id": "manual", "broker": manual_broker or broker, "name": "사용자 별칭",
            "owner": "아빠", "account_no": account_no,
        }])
        written = {}
        with patch.object(main, provider_attr, return_value=provider), \
             patch.object(main, "read_portfolio", return_value=deepcopy(data)), \
             patch.object(main, "write_portfolio", side_effect=lambda value, **_: written.update(value)):
            await runner("reconciliation-fixture")
        self.assertEqual(len(written["accounts"]), 1)
        self.assertEqual(written["accounts"][0]["id"], "manual")
        self.assertEqual(written["accounts"][0]["name"], "사용자 별칭")
        self.assertEqual(written["accounts"][0]["owner"], "아빠")

    async def test_nh_reconciliation_reuses_manual_alias_and_account_number(self):
        class NH:
            configured = True
            last_accounts = [{"acct_no": "12345678", "acct_type": "01"}]
            account_cash = {}
            async def sync_holdings(self):
                return BrokerHoldingsResult.authoritative_result([], (ProviderHoldingScope("12345678", ALL_MARKETS),), cash_valid=True)
            def _account_name(self, _): return "NH API"
        await self._assert_provider_reuses_manual_account(
            provider_attr="NhPlugOpenAPI", provider=NH(), runner=main.sync_namoo_for_user,
            account_no="1234-5678", broker="NH투자증권(나무)", manual_broker="나무증권",
        )

    async def test_kis_reconciliation_normalizes_product_code_format(self):
        class KIS:
            configured = True
            last_accounts = [{"account_number": "1234567801", "account_name": "KIS API"}]
            account_cash = {}
            def _parse_account_no(self): return "12345678", "01"
            async def sync_holdings(self):
                return BrokerHoldingsResult.authoritative_result([], (ProviderHoldingScope("1234567801", ALL_MARKETS),), cash_valid=True)
        await self._assert_provider_reuses_manual_account(
            provider_attr="KISOpenAPI", provider=KIS(), runner=main.sync_kis_for_user,
            account_no="12345678-01", broker="한국투자증권", manual_broker="KIS",
        )

    async def test_kiwoom_reconciliation_reuses_manual_account(self):
        class Kiwoom:
            configured = True
            last_accounts = [{"account_number": "12345678", "account_name": "Kiwoom API"}]
            account_cash = {}
            async def sync_holdings(self):
                return BrokerHoldingsResult.authoritative_result([], (ProviderHoldingScope("12345678", ALL_MARKETS),), cash_valid=True)
        await self._assert_provider_reuses_manual_account(
            provider_attr="KiwoomOpenAPI", provider=Kiwoom(), runner=main.sync_kiwoom_for_user,
            account_no="1234-5678", broker="키움증권", manual_broker="키움",
        )

    async def test_toss_full_manual_number_reuses_only_unambiguous_suffix(self):
        data = self._empty_portfolio(accounts=[{
            "id": "manual", "broker": "토스", "name": "사용자 별칭", "owner": "아빠",
            "account_no": "1234567890",
        }])
        class Toss:
            configured = True
            last_accounts = [{"accountSeq": 7, "accountNo": "0000007890"}]
            async def sync_holdings(self):
                return BrokerHoldingsResult.authoritative_result([], (ProviderHoldingScope("7", ALL_MARKETS),), cash_valid=True)
            async def get_buying_power(self, _): return {"KRW": 0}
        written = {}
        with patch.object(main, "TossOpenAPI", return_value=Toss()), \
             patch.object(main, "read_portfolio", return_value=deepcopy(data)), \
             patch.object(main, "write_portfolio", side_effect=lambda value, **_: written.update(value)):
            await main.sync_toss_for_user("reconciliation-fixture")
        self.assertEqual(len(written["accounts"]), 1)
        self.assertEqual(written["accounts"][0]["id"], "manual")
        self.assertEqual(written["accounts"][0]["name"], "사용자 별칭")
        self.assertEqual(written["accounts"][0]["account_no"], "7890")
